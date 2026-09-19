import io
import pytest
import fitz  # PyMuPDF
from docx import Document
from fastapi import HTTPException
from app.services.resume_parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_resume_document
)
from app.services.text_cleaner import clean_extracted_text, estimate_word_count


def test_text_cleaner():
    """Verify unicode bullets, excessive line breaks, and whitespace are normalized"""
    raw = "Jane Doe \t\t\n\n\n\uf0b7 Python Programming\n\u2022 FastAPI Development\r\n\r\n\n- SQL Database"
    cleaned = clean_extracted_text(raw)
    assert "Jane Doe" in cleaned
    assert "Python Programming" in cleaned
    assert "FastAPI Development" in cleaned
    assert "\uf0b7" not in cleaned
    assert "\u2022" not in cleaned
    assert estimate_word_count(cleaned) > 0


def test_pdf_text_extraction():
    """Verify in-memory PDF generation and text extraction via PyMuPDF"""
    doc = fitz.open()
    page = doc.new_page()
    sample_text = (
        "Jane Doe\n"
        "Full Stack Python & React Engineer\n"
        "Skills: Python, FastAPI, React, PostgreSQL, Docker, Git\n"
        "Experience: 2 years as Software Engineer building APIs\n"
        "Education: BSc in Information Technology"
    )
    page.insert_text((50, 72), sample_text)
    pdf_bytes = doc.tobytes()

    extracted = extract_text_from_pdf(pdf_bytes)
    assert "Jane Doe" in extracted
    assert "Python" in extracted
    assert "FastAPI" in extracted


def test_docx_text_extraction():
    """Verify in-memory DOCX generation, paragraphs and tables extraction via python-docx"""
    doc = Document()
    doc.add_heading("John Smith - Resume", level=1)
    doc.add_paragraph("Summary: Experienced Data Analyst with strong SQL and Python background.")
    
    # Add a table to test table cell extraction
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Core Competencies"
    table.rows[0].cells[1].text = "Python, Pandas, PowerBI, Advanced Excel"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    extracted = extract_text_from_docx(docx_bytes)
    assert "John Smith" in extracted
    assert "Data Analyst" in extracted
    assert "PowerBI" in extracted


def test_scanned_pdf_detection():
    """Verify that a PDF without readable text raises 422 Unprocessable Entity"""
    # Create empty PDF page without any text
    doc = fitz.open()
    doc.new_page()
    empty_pdf_bytes = doc.tobytes()

    with pytest.raises(HTTPException) as exc_info:
        extract_text_from_pdf(empty_pdf_bytes)
    assert exc_info.value.status_code == 422
    assert "scanned" in exc_info.value.detail.lower()


def test_unified_parse_resume_document():
    """Verify parse_resume_document correctly detects format and cleans text"""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Alex Johnson\nEmail: alex@example.com\nSkills: Python, Machine Learning, TensorFlow, Pandas")
    pdf_bytes = doc.tobytes()

    result = parse_resume_document(file_bytes=pdf_bytes, filename="alex_resume.pdf")
    assert "Alex Johnson" in result
    assert "TensorFlow" in result
